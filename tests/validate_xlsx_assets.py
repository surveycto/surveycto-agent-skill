#!/usr/bin/env python3
"""Validate public XLSX assets for metadata leaks and template invariants."""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = REPO_ROOT / "assets"
XLSFORM_TEMPLATE = ASSET_DIR / "xlsform-template.xlsx"
ALLOWED_AUTHORS = {"SurveyCTO", "Dobility, Inc. (SurveyCTO)"}
STRUCTURAL_PATH_PARTS = {"xl/workbook.xml", "docProps/app.xml", "docProps/custom.xml"}
LOCAL_PATH_PATTERNS = [
    re.compile(r"/Users/[^\s<>'\"]+"),
    re.compile(r"/home/[^\s<>'\"]+"),
    re.compile(r"[A-Za-z]:\\Users\\[^\s<>'\"]+"),
]
CORE_NS = {
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
}
EXPECTED_TEMPLATE_SHEETS = [
    "survey",
    "choices",
    "settings",
    "help-survey",
    "help-choices",
    "help-settings",
]
EXPECTED_COLUMNS = {
    "survey": [
        "type",
        "name",
        "label",
        "hint",
        "default",
        "appearance",
        "constraint",
        "constraint message",
        "relevance",
        "disabled",
        "required",
        "required message",
        "read only",
        "calculation",
        "repeat_count",
        "media:image",
        "media:audio",
        "media:video",
        "choice_filter",
        "note",
        "response_note",
        "publishable",
        "minimum_seconds",
    ],
    "choices": ["list_name", "value", "label", "image", "filter"],
    "settings": [
        "form_title",
        "form_id",
        "version",
        "public_key",
        "submission_url",
        "default_language",
    ],
}
EXPECTED_VERSION_FORMULA = (
    '=TEXT(YEAR(NOW())-2000, "00") & TEXT(MONTH(NOW()), "00") & '
    'TEXT(DAY(NOW()), "00") & TEXT(HOUR(NOW()), "00") & '
    'TEXT(MINUTE(NOW()), "00")'
)


def validate_workbook(workbook: Path) -> list[str]:
    errors: list[str] = []

    with zipfile.ZipFile(workbook) as archive:
        if "docProps/core.xml" in archive.namelist():
            core = archive.read("docProps/core.xml")
            root = ElementTree.fromstring(core)
            creator = root.findtext("dc:creator", namespaces=CORE_NS)
            if creator and creator not in ALLOWED_AUTHORS:
                errors.append(f"{workbook}: docProps/core.xml has unexpected creator={creator!r}")

            last_modified_by = root.findtext("cp:lastModifiedBy", namespaces=CORE_NS)
            if last_modified_by and last_modified_by not in ALLOWED_AUTHORS:
                errors.append(
                    f"{workbook}: docProps/core.xml has unexpected lastModifiedBy={last_modified_by!r}"
                )

        for part in [name for name in archive.namelist() if name.endswith(".xml")]:
            text = archive.read(part).decode("utf-8", errors="replace")
            if "x15ac:absPath" in text:
                errors.append(f"{workbook}: {part} contains x15ac:absPath")

            if part not in STRUCTURAL_PATH_PARTS:
                continue

            for pattern in LOCAL_PATH_PATTERNS:
                match = pattern.search(text)
                if match:
                    errors.append(f"{workbook}: {part} contains local path {match.group(0)!r}")

    return errors


def header_values(sheet) -> list[str]:
    values: list[str] = []
    for cell in sheet[1]:
        if cell.value is not None:
            values.append(cell.value)
    return values


def validate_xlsform_template(workbook: Path) -> list[str]:
    errors: list[str] = []
    wb = load_workbook(workbook, data_only=False)

    if wb.sheetnames != EXPECTED_TEMPLATE_SHEETS:
        errors.append(f"{workbook}: sheet structure is {wb.sheetnames!r}")

    for sheet_name, expected_columns in EXPECTED_COLUMNS.items():
        if sheet_name not in wb.sheetnames:
            continue

        actual_columns = header_values(wb[sheet_name])
        if actual_columns != expected_columns:
            errors.append(
                f"{workbook}: {sheet_name} headers are {actual_columns!r}; "
                f"expected {expected_columns!r}"
            )

    if "survey" in wb.sheetnames:
        survey = wb["survey"]
        expected_columns = EXPECTED_COLUMNS["survey"]
        covered_columns: set[int] = set()

        for conditional_formatting in survey.conditional_formatting:
            if not conditional_formatting.rules:
                continue

            for cell_range in conditional_formatting.sqref.ranges:
                if cell_range.max_row < 2:
                    continue
                start = max(1, cell_range.min_col)
                end = min(len(expected_columns), cell_range.max_col)
                covered_columns.update(range(start, end + 1))

        missing_columns = [
            get_column_letter(column)
            for column in range(1, len(expected_columns) + 1)
            if column not in covered_columns
        ]
        if missing_columns:
            errors.append(
                f"{workbook}: survey conditional formatting does not cover "
                f"headered columns {missing_columns!r}"
            )

    if "settings" in wb.sheetnames:
        actual_formula = wb["settings"]["C2"].value
        if actual_formula != EXPECTED_VERSION_FORMULA:
            errors.append(
                f"{workbook}: settings!C2 formula is {actual_formula!r}; "
                f"expected {EXPECTED_VERSION_FORMULA!r}"
            )

    return errors


def main() -> int:
    errors: list[str] = []
    for workbook in sorted(ASSET_DIR.rglob("*.xlsx")):
        errors.extend(validate_workbook(workbook))

    if XLSFORM_TEMPLATE.exists():
        errors.extend(validate_xlsform_template(XLSFORM_TEMPLATE))
    else:
        errors.append(f"{XLSFORM_TEMPLATE}: missing XLSForm template")

    if errors:
        print("XLSX asset validation failed:", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1

    print("XLSX asset validation passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
